import io
from pathlib import Path
from typing import Any
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Inches as _in, Pt
from backend.app.services.layout_engine import SlideLayoutEngine
SLIDE_WIDTH, SLIDE_HEIGHT = _in(13.33), _in(7.5)
CT, CX, CH = RGBColor(26, 26, 46), RGBColor(45, 45, 45), RGBColor(15, 60, 120)
CM = {"line": XL_CHART_TYPE.LINE, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "pie": XL_CHART_TYPE.PIE, "area": XL_CHART_TYPE.AREA, "doughnut": XL_CHART_TYPE.DOUGHNUT}
_X = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
_LE = SlideLayoutEngine()
def _tr(slide: Any, t: str = "fade", s: str = "med") -> None:
    t = t if t in {"fade", "push", "wipe", "cut", "zoom", "cover", "split"} else "fade"
    slide._element.append(parse_xml(f'<p:transition {_X} spd="{s}" advClick="1"><p:{t}/></p:transition>'))
def _an(slide: Any, shape_ids: list[int] | None = None, e: str = "appear") -> None:
    slide._element.append(parse_xml(f'<p:timing {_X}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="always" nodeType="tmRoot"/></p:par></p:tnLst></p:timing>'))
_apply_slide_transition = _tr
_apply_slide_animations = _an
def _tb(slide: Any, text: str, l: Any, t: Any, w: Any, h: Any, fs: int = 18, b: bool = False, c: RGBColor = CX, a: Any = PP_ALIGN.LEFT) -> Any:
    tx = slide.shapes.add_textbox(l, t, w, h); p = tx.text_frame.paragraphs[0]; p.alignment = a
    run = p.add_run(); run.text = text; run.font.size, run.font.bold, run.font.color.rgb = Pt(fs), b, c
    return tx
def _sn(slide: Any, sd: dict[str, Any]) -> None:
    if notes := sd.get("speaker_notes", ""): slide.notes_slide.notes_text_frame.text = notes
def _ms(prs: Presentation, li: int, sd: dict[str, Any]) -> Any:
    s = prs.slides.add_slide(prs.slide_layouts[li])
    if title := sd.get("title", ""):
        if s.shapes.title: s.shapes.title.text = title
        else: _tb(s, title, _in(1.0), _in(0.5), _in(12.0), _in(1.0), 24, True, CT)
    return s
def _fin(s: Any, d: dict[str, Any], tr: str = "wipe", an: str = "wipe", shape_ids: list[int] | None = None) -> None:
    _sn(s, d); _tr(s, tr); _an(s, shape_ids, an)
def _cts(prs: Presentation, d: dict[str, Any]) -> None:
    _fin(_ms(prs, 0, d), d, "fade", "fade")
def _ccs(prs: Presentation, d: dict[str, Any]) -> None:
    lay, s, body = _LE.compute_slide_layout(d), _ms(prs, 1, d), d.get("body_text", "")
    if lay.columns.count > 1 and body:
        lines = [l.strip() for l in body.split("\n") if l.strip()]; pc = max(1, -(-len(lines) // lay.columns.count))
        for ci, cb in enumerate(lay.columns.column_boxes):
            if cl := lines[ci * pc:(ci + 1) * pc]: _tb(s, "\n".join(cl), _in(cb.left), _in(cb.top), _in(cb.width), _in(cb.height), lay.typography.body_pt, False, CX)
    elif len(s.shapes.placeholders) > 1 and body: s.shapes.placeholders[1].text = body
    _fin(s, d)
def _cch(prs: Presentation, d: dict[str, Any]) -> None:
    cb, s, cd = _LE.compute_slide_layout(d).content_box, _ms(prs, 5, d), d.get("chart_data") or {}
    ct = CM.get(str(cd.get("chart_type", "bar")).lower(), XL_CHART_TYPE.BAR_CLUSTERED)
    cdata = ChartData(); cdata.categories = cd.get("categories", ["Cat 1", "Cat 2"])
    for series in cd.get("series", []): cdata.add_series(series.get("name", "Series"), tuple(series.get("values", [])))
    if not cd.get("series") and cd.get("data"): cdata.add_series("Series", tuple(v for _, v in cd.get("data", [])))
    cs = s.shapes.add_chart(ct, _in(cb.left), _in(cb.top), _in(cb.width), _in(cb.height), cdata)
    _fin(s, d, "zoom", "zoom", [cs.shape_id])
def _ctb(prs: Presentation, d: dict[str, Any]) -> None:
    cb, s, td = _LE.compute_slide_layout(d).content_box, _ms(prs, 5, d), d.get("table_data") or {}
    hdr, rows = td.get("headers", []), td.get("rows", [])
    ts = s.shapes.add_table(len(rows) + (1 if hdr else 0), max(1, len(hdr)), _in(cb.left), _in(cb.top), _in(cb.width), _in(cb.height))
    for ci, h in enumerate(hdr): ts.table.cell(0, ci).text = str(h)
    for ri, r in enumerate(rows, start=1 if hdr else 0):
        for ci, val in enumerate(r): ts.table.cell(ri, ci).text = str(val if val is not None else "")
    _fin(s, d, "wipe", "wipe", [ts.shape_id])
def _cks(prs: Presentation, d: dict[str, Any]) -> None:
    lay, s = _LE.compute_slide_layout(d), _ms(prs, 5, d)
    bullets = d.get("bullets") or d.get("kpi", {}).get("bullets", [])
    if len(s.shapes.placeholders) > 1 and d.get("kpi", {}).get("label"): s.shapes.placeholders[1].text = d["kpi"]["label"]
    ih, gap, pc = lay.extras.get("item_height_in", 0.7), lay.extras.get("row_gap_in", 0.15), max(1, -(-len(bullets) // max(1, lay.columns.count)))
    for ci, col_box in enumerate(lay.columns.column_boxes):
        for ri, b in enumerate(bullets[ci * pc:(ci + 1) * pc]):
            _tb(s, b, _in(col_box.left), _in(col_box.top + ri * (ih + gap)), _in(col_box.width), _in(ih), lay.typography.body_pt, True, CT)
    _fin(s, d, "appear", "appear")
def _ccp(prs: Presentation, d: dict[str, Any]) -> None:
    lay, s, cmp = _LE.compute_slide_layout(d), _ms(prs, 5, d), d.get("comparison") or {}
    if lay.columns.count == 2 and len(lay.columns.column_boxes) == 2:
        lb, rb, bpt = lay.columns.column_boxes[0], lay.columns.column_boxes[1], lay.typography.body_pt
        for pfx, box, clr in (("left", lb, CT), ("right", rb, CH)):
            if t := cmp.get(f"{pfx}_title"): _tb(s, t, _in(box.left), _in(box.top), _in(box.width), _in(0.5), bpt + 2, True, clr)
            if b := cmp.get(f"{pfx}_body"): _tb(s, b, _in(box.left), _in(box.top + 0.6), _in(box.width), _in(box.height - 0.6), bpt, False, CX)
    elif len(s.shapes.placeholders) > 1 and cmp.get("left_title"): s.shapes.placeholders[1].text = str(cmp["left_title"])
    _fin(s, d)
def _cps(prs: Presentation, d: dict[str, Any]) -> None:
    s = _ms(prs, 1, d)
    if len(s.shapes.placeholders) > 1 and d.get("process", {}).get("steps"): s.shapes.placeholders[1].text = str(d["process"]["steps"])
    _fin(s, d, "push", "push")
_COMPILERS = {"title": _cts, "content": _ccs, "chart": _cch, "table": _ctb, "kpi": _cks, "comparison": _ccp, "process": _cps}
class PptxCompiler:
    def compile_deck(self, deck_title: str, slides: list[dict[str, Any]], template_bytes: bytes | None = None, template_path: str | None = None) -> bytes:
        if template_bytes: prs = Presentation(io.BytesIO(template_bytes))
        elif template_path and Path(template_path).is_file(): prs = Presentation(template_path)
        else: prs = Presentation(); prs.slide_width, prs.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT
        for s in slides: _COMPILERS.get(s.get("slide_type", "content"), _ccs)(prs, s)
        out = io.BytesIO(); prs.save(out); return out.getvalue()
import argparse
import base64
import json
import sys


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="PPTX Compiler CLI")
    p.add_argument("--input", required=True, help="JSON file with deck data: {\"title\": str, \"slides\": [...], \"template\": str?}")
    p.add_argument("--output", required=True, help="Output PPTX file path")
    p.add_argument("--json", action="store_true", help="Output base64-encoded PPTX as JSON to stdout")
    args = p.parse_args(argv)

    try:
        with open(args.input, "r") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        sys.stderr.write(f"Failed to read input: {e}\n")
        return 1

    title = data.get("title", "Deck")
    slides = data.get("slides", [])
    template_path = data.get("template")

    compiler = PptxCompiler()
    try:
        pptx_bytes = compiler.compile_deck(title, slides, template_path=template_path)
    except Exception as e:
        sys.stderr.write(f"Compilation failed: {e}\n")
        return 1

    if args.json:
        sys.stdout.write(json.dumps({"pptx_base64": base64.b64encode(pptx_bytes).decode("ascii"), "output": args.output}) + "\n")
    else:
        try:
            with open(args.output, "wb") as f:
                f.write(pptx_bytes)
            sys.stdout.write(f"Written {len(pptx_bytes)} bytes to {args.output}\n")
        except OSError as e:
            sys.stderr.write(f"Failed to write output: {e}\n")
            return 1

    return 0


if __name__ == "__main__":
    sys.exit(main())
