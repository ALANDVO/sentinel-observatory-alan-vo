from __future__ import annotations
import io, os
from pathlib import Path
from sqlalchemy.orm import Session
from backend.app.models.template import Template
class TemplateService:
    TEMPLATES_DIR = os.getenv("TEMPLATES_DIR", "./data/templates")
    def create_template(self, db: Session, name: str, description: str | None, author_sub: str | None, pptx_bytes: bytes | None = None) -> Template:
        layout_names, placeholder_map, slide_count, filename, theme_color = [], {}, 0, None, None
        if pptx_bytes:
            layout_names, placeholder_map, slide_count, theme_color, filename = self._scan_pptx(name, pptx_bytes)
        t = Template(name=name, description=description or "", author_sub=author_sub, filename=filename,
                     theme_color=theme_color, slide_count=slide_count, layout_names=layout_names,
                     placeholder_map=placeholder_map, is_builtin="false")
        db.add(t); db.commit(); db.refresh(t)
        return t
    def _scan_pptx(self, name: str, pptx_bytes: bytes) -> tuple[list[str], dict[str, list], int, str | None, str | None]:
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(pptx_bytes))
            layout_names = [l.name for l in prs.slide_layouts]
            placeholder_map = {l.name: [{"idx": ph.placeholder_format.idx, "type": str(ph.placeholder_format.type), "name": ph.name} for ph in l.placeholders] for l in prs.slide_layouts}
            safe_name = "".join(c if c.isalnum() or c in "-_" else "_" for c in name)
            os.makedirs(self.TEMPLATES_DIR, exist_ok=True)
            filename = f"{safe_name}.pptx"
            with open(Path(self.TEMPLATES_DIR) / filename, "wb") as f: f.write(pptx_bytes)
            return layout_names, placeholder_map, len(prs.slides), None, filename
        except Exception as exc:
            raise ValueError(f"Failed to parse PPTX template: {exc}") from exc
    def list_templates(self, db: Session, skip: int = 0, limit: int = 50) -> tuple[list[Template], int]:
        return db.query(Template).offset(skip).limit(limit).all(), db.query(Template).count()
    def get_template(self, db: Session, template_id: str) -> Template | None:
        return db.query(Template).filter(Template.id == template_id).first()
    def get_template_filepath(self, template: Template) -> Path | None:
        if not template or not template.filename: return None
        p = Path(self.TEMPLATES_DIR) / template.filename
        return p if p.is_file() else None
    def delete_template(self, db: Session, template_id: str) -> bool:
        if not (t := self.get_template(db, template_id)): return False
        if t.filename: Path(self.TEMPLATES_DIR, t.filename).unlink(missing_ok=True)
        db.delete(t); db.commit()
        return True
    def seed_builtin_templates(self, db: Session) -> None:
        if db.query(Template).filter(Template.is_builtin == "true").count() > 0: return
        for n, d, c, l in [
            ("Executive Brief", "Minimal executive deck.", "#1A1A2E", ["Title Slide", "Content", "Chart", "Table"]),
            ("Research Report", "Academic research findings.", "#0F3C78", ["Title Slide", "Section Header", "Comparison"]),
            ("Data Dashboard", "Data-dense KPI template.", "#162136", ["Title Slide", "KPI Overview", "Chart", "Table"]),
        ]:
            db.add(Template(name=n, description=d, theme_color=c, layout_names=l, placeholder_map={}, is_builtin="true"))
        db.commit()
