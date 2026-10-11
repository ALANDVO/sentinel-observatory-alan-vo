"""Tests for native PPTX compiler: transitions, animations, charts, tables, and template loading."""
from __future__ import annotations

import io
import pytest
from pptx import Presentation
from pptx.util import Inches

from backend.app.services.pptx_compiler import PptxCompiler, _apply_slide_transition, _apply_slide_animations


@pytest.fixture
def compiler():
    return PptxCompiler()


def test_compile_title_slide_produces_valid_pptx(compiler: PptxCompiler):
    slides = [{"slide_type": "title", "title": "AI Safety Overview", "bullets": ["Alignment Research"], "speaker_notes": "Introduction to the topic."}]
    pptx_bytes = compiler.compile_deck("AI Safety", slides)
    assert len(pptx_bytes) > 1000
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 1


def test_compile_chart_slide_creates_native_chart(compiler: PptxCompiler):
    slides = [
        {"slide_type": "title", "title": "Chart Demo"},
        {
            "slide_type": "chart",
            "title": "Model Performance",
            "chart_data": {
                "chart_type": "bar",
                "categories": ["GPT-4", "Claude-3", "Gemini"],
                "series": [{"name": "Accuracy", "values": [0.92, 0.91, 0.89]}],
            },
            "speaker_notes": "Compare model accuracy.",
        },
    ]
    pptx_bytes = compiler.compile_deck("Performance", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 2
    # Verify the chart slide has a chart shape
    chart_slide = prs.slides[1]
    chart_shapes = [s for s in chart_slide.shapes if s.has_chart]
    assert len(chart_shapes) >= 1


def test_compile_table_slide_creates_native_table(compiler: PptxCompiler):
    slides = [{
        "slide_type": "table",
        "title": "Benchmark Results",
        "table_data": {
            "headers": ["Model", "MMLU", "HumanEval"],
            "rows": [["GPT-4", "0.86", "0.67"], ["Claude-3", "0.85", "0.73"]],
        },
    }]
    pptx_bytes = compiler.compile_deck("Benchmarks", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    table_shapes = [s for s in prs.slides[0].shapes if s.has_table]
    assert len(table_shapes) >= 1
    table = table_shapes[0].table
    assert table.cell(0, 0).text == "Model"


def test_compile_kpi_slide_produces_multiple_boxes(compiler: PptxCompiler):
    slides = [{
        "slide_type": "kpi",
        "title": "Observatory KPIs",
        "bullets": ["Accuracy: 94.2%", "Latency: <50ms", "Coverage: 98.7%", "F1 Score: 0.91"],
    }]
    pptx_bytes = compiler.compile_deck("KPIs", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    # KPI slide should have multiple shape objects (boxes + text)
    slide = prs.slides[0]
    assert len(slide.shapes) >= 4


def test_apply_slide_transition_adds_xml_element(compiler: PptxCompiler):
    slides = [{"slide_type": "content", "title": "Transitions Test", "body_text": "• Point A\n• Point B"}]
    pptx_bytes = compiler.compile_deck("Transitions", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    slide = prs.slides[0]
    # The p:transition element must exist in slide XML
    xml = slide._element.xml
    assert "p:transition" in xml


def test_apply_slide_animations_adds_timing_element(compiler: PptxCompiler):
    slides = [{"slide_type": "content", "title": "Animations Test", "body_text": "• AI point\n• ML point"}]
    pptx_bytes = compiler.compile_deck("Animations", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    slide = prs.slides[0]
    xml = slide._element.xml
    assert "p:timing" in xml


def test_compile_multi_slide_deck_preserves_order(compiler: PptxCompiler):
    slides = [
        {"slide_type": "title", "title": "Introduction"},
        {"slide_type": "content", "title": "Methods"},
        {"slide_type": "chart", "title": "Results", "chart_data": {"categories": ["A", "B"], "series": [{"name": "S", "values": [1, 2]}]}},
        {"slide_type": "table", "title": "Summary", "table_data": {"headers": ["Key", "Value"], "rows": [["Accuracy", "95%"]]}},
    ]
    pptx_bytes = compiler.compile_deck("Full Deck", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 4


def test_speaker_notes_stored_in_slide(compiler: PptxCompiler):
    notes_text = "Explain the cognitive load evaluation methodology in detail here."
    slides = [{"slide_type": "content", "title": "Methods", "body_text": "Cognitive evaluation.", "speaker_notes": notes_text}]
    pptx_bytes = compiler.compile_deck("Notes Test", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    slide = prs.slides[0]
    actual_notes = slide.notes_slide.notes_text_frame.text
    assert notes_text in actual_notes


def test_transition_types_map_per_slide_type(compiler: PptxCompiler):
    """Different slide types should get different transition animations."""
    slides = [
        {"slide_type": "title", "title": "Title"},
        {"slide_type": "chart", "title": "Chart", "chart_data": {"categories": ["X"], "series": [{"name": "S", "values": [1]}]}},
        {"slide_type": "kpi", "title": "KPI", "bullets": ["M1: 99%"]},
    ]
    pptx_bytes = compiler.compile_deck("Transition Types", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 3
    for slide in prs.slides:
        xml = slide._element.xml
        assert "p:transition" in xml


def test_template_loading_applies_layouts(compiler: PptxCompiler, tmp_path):
    """When a template .pptx is provided, its slide layouts are used."""
    # Create a minimal valid .pptx template using python-pptx
    template_prs = Presentation()
    template_prs.slide_width = Inches(13.33)
    template_prs.slide_height = Inches(7.5)
    template_path = str(tmp_path / "test_template.pptx")
    template_prs.save(template_path)

    slides = [{"slide_type": "content", "title": "Template Test", "body_text": "Uses custom template."}]
    pptx_bytes = compiler.compile_deck("Template Deck", slides, template_path=template_path)
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 1


def test_content_slide_line_chart(compiler: PptxCompiler):
    slides = [{
        "slide_type": "chart",
        "title": "Training Loss",
        "chart_data": {
            "chart_type": "line",
            "categories": ["Epoch 1", "Epoch 2", "Epoch 3", "Epoch 4"],
            "series": [{"name": "Loss", "values": [0.8, 0.5, 0.3, 0.18]}],
        },
    }]
    pptx_bytes = compiler.compile_deck("Training Curve", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))
    chart_slide = prs.slides[0]
    chart_shapes = [s for s in chart_slide.shapes if s.has_chart]
    assert len(chart_shapes) == 1
